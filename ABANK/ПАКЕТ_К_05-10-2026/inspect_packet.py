from pathlib import Path
import json,re
import pypdfium2 as pdfium
from docx import Document
from zipfile import ZipFile
from xml.etree import ElementTree as ET

R=Path(__file__).resolve().parent;Q=R/'qa';Q.mkdir(exist_ok=True)
files=json.loads((R/'packet_files.json').read_text(encoding='utf-8'))
report=[]
for f in files:
    p=R/f['docx'];pdf=p.with_suffix('.pdf')
    doc=Document(p)
    text='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    if f['number']<5:
        for t in ['127/20805/26','2/127/5742/26','Подольський Андрій Юрійович','Адреса реєстрації']:
            assert t in text,(f['number'],t)
        for t in ['фактично проживаю','перебуваю в іншому','фактична адреса','місце фактичного проживання']:
            assert t not in text.lower(),(f['number'],t)
    with ZipFile(p) as z:
        styles=ET.fromstring(z.read('word/styles.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for st in styles.findall('w:style',ns):
            if st.get('{'+ns['w']+'}styleId') in ['Title','Subtitle','Heading1','Heading2']:
                assert st.find('w:pPr/w:pBdr',ns) is None,(f['number'],'title border')
    with pdfium.PdfDocument(pdf) as d:
        pageinfo=[];chunks=[]
        for n,page in enumerate(d,1):
            tp=page.get_textpage();t=tp.get_text_range();chunks.append(t)
            dest=Q/f"{f['number']:02}-page-{n:02}.png"
            page.render(scale=1.5).to_pil().save(dest)
            assert len(t.strip())>10,(f['number'],n,'empty page')
            pageinfo.append({'page':n,'chars':len(t),'image':str(dest.relative_to(R)),'start':' '.join(t.split())[:95],'end':' '.join(t.split())[-95:]})
            tp.close();page.close()
        full='\n'.join(chunks)
        assert '\ufffd' not in full,(f['number'],'replacement glyph')
        (Q/f"{f['number']:02}-rendered.txt").write_text(full,encoding='utf-8')
        for token in ['74171/26-Вх'] if f['number']==1 else []:assert token in full
        report.append({'number':f['number'],'pages':len(d),'pageinfo':pageinfo})
        f['pdf']=str(pdf.relative_to(R));f['pages']=len(d)
(R/'packet_files.json').write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
(Q/'render_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
